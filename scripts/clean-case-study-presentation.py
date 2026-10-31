#!/usr/bin/env python3
"""Update repeated PPT footer and render nine preview images from the edited deck."""
import base64, json, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
import fitz
root=Path(__file__).resolve().parents[1]
ppt=root/"src/case-study/downloads/footmate_report.pptx"
preview=root/"src/case-study/report-previews.js"
pres=Presentation(str(ppt))
assert len(pres.slides)==9
footer="가정·합성 데이터 기반 / 실제 사업 실적 아님"
specific={2:{"전체 정원과 수익 대상 거래를 구분":"기준 가정 · 정원과 경유 예약 구분"},4:{"상반기 합성 실적과 하반기 계획을 연결한 예측":"상반기 모형 거래·하반기 전망 비교"},5:{"6월 차이":"6월 모형 차이"}}
counts={}
def edit(tf, replacements):
    for p in tf.paragraphs:
        for a,b in replacements.items():
            if a not in p.text:continue
            direct=False
            for run in p.runs:
                if a in run.text:
                    run.text=run.text.replace(a,b);direct=True;break
            if not direct:
                if not p.runs:raise ValueError("No run for "+a)
                p.runs[0].text=p.text.replace(a,b)
                for run in p.runs[1:]:run.text=""
            counts[a]=counts.get(a,0)+1
for i,slide in enumerate(pres.slides):
    repl={footer:"사업·IT 기획",**specific.get(i,{})}
    for shape in slide.shapes:
        if shape.has_text_frame:edit(shape.text_frame,repl)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:edit(cell.text_frame,repl)
assert counts.get(footer)==9,counts
for i in specific:
    for k in specific[i]:assert counts.get(k)==1,counts
pres.save(str(ppt))
assert len(Presentation(str(ppt)).slides)==9
start="window.REPORT_SLIDES="
raw=preview.read_text(encoding="utf-8").strip()
assert raw.startswith(start)
images=json.loads(raw[len(start):].rstrip(";"))
assert len(images)==9
with tempfile.TemporaryDirectory() as d:
    subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/footmate-copy-clean-lo","--headless","--convert-to","pdf","--outdir",d,str(ppt)],check=True,timeout=240)
    pdf=fitz.open(str(Path(d)/"footmate_report.pdf"))
    assert len(pdf)==9
    for i,page in enumerate(pdf):
        png=page.get_pixmap(matrix=fitz.Matrix(4/3,4/3),alpha=False).tobytes("png")
        images[i]["src"]="data:image/png;base64,"+base64.b64encode(png).decode("ascii")
preview.write_text(start+json.dumps(images,ensure_ascii=False,separators=(",",":"))+";\n",encoding="utf-8")
print("PPT_UPDATED",counts,"PPT_BYTES",ppt.stat().st_size,"PREVIEW_BYTES",preview.stat().st_size)
